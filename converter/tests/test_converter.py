"""Converter tests: each conversion on real files, the refusals, the LibreOffice hardening, and the HTTP service's
access and limit checks. Needs LibreOffice (`soffice`) and requirements.txt. Run: python -m unittest -v
"""

from __future__ import annotations

import http.client
import http.server
import io
import json
import os
import shutil
import subprocess
import tempfile
import threading
import unittest
from unittest.mock import patch
import zipfile

from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation

from rokela_convert import engines, server
from rokela_convert.engines import UserError

TOKEN = "t" * 40
HAS_SOFFICE = shutil.which("soffice") is not None


def soffice_pdf(src: str, outdir: str, filter_opts: str = "") -> str:
    """Makes a fixture PDF with plain LibreOffice (test setup only)."""
    prof = tempfile.mkdtemp()
    target = "pdf" + (f":writer_pdf_Export:{filter_opts}" if filter_opts else "")
    subprocess.run(
        ["soffice", f"-env:UserInstallation=file://{prof}", "--headless", "--convert-to", target, "--outdir", outdir, src],
        check=True,
        capture_output=True,
        timeout=120,
    )
    shutil.rmtree(prof, ignore_errors=True)
    return os.path.join(outdir, os.path.splitext(os.path.basename(src))[0] + ".pdf")


def report_docx(path: str, pages: int = 1) -> None:
    d = Document()
    d.add_heading("Quarterly report", 0)
    p = d.add_paragraph("The first paragraph has some ")
    p.add_run("bold words").bold = True
    p.add_run(" and then carries on long enough to wrap across more than one line, so the lines must be joined again.")
    d.add_paragraph('A cell-like line =HYPERLINK("http://example.invalid") stays text.')
    t = d.add_table(rows=3, cols=3)
    t.style = "Table Grid"
    for r, row in enumerate([["Item", "Qty", "Price"], ["Apples", "12", "1,250.50"], ["=cmd|' /C calc'!A0", "3", "7"]]):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    for i in range(1, pages):
        d.add_page_break()
        d.add_paragraph(f"Page {i + 1} text")
    d.save(path)


@unittest.skipUnless(HAS_SOFFICE, "LibreOffice is not installed")
class CellValueTests(unittest.TestCase):
    def test_numbers_are_typed_but_identifiers_keep_every_digit(self) -> None:
        self.assertEqual(engines._cell_value("42"), 42)
        self.assertEqual(engines._cell_value("1,234.50"), 1234.5)
        self.assertEqual(engines._cell_value("0"), 0)
        self.assertEqual(engines._cell_value("0.75"), 0.75)
        self.assertEqual(engines._cell_value("-12"), -12)
        self.assertEqual(engines._cell_value("123456789012345"), 123456789012345)
        for ident in ["12345678901234567", "1234567890123456", "007", "0761234567", "1234567890.1234567"]:
            self.assertEqual(engines._cell_value(ident), ident)


class EngineTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.dir = tempfile.mkdtemp()
        report_docx(os.path.join(cls.dir, "report.docx"), pages=2)
        cls.pdf = soffice_pdf(os.path.join(cls.dir, "report.docx"), cls.dir)
        opts = '{"EncryptFile":{"type":"boolean","value":"true"},"DocumentOpenPassword":{"type":"string","value":"secret1"}}'
        enc_dir = os.path.join(cls.dir, "enc")
        os.makedirs(enc_dir)
        cls.encrypted = soffice_pdf(os.path.join(cls.dir, "report.docx"), enc_dir, opts)

    @classmethod
    def tearDownClass(cls) -> None:
        shutil.rmtree(cls.dir, ignore_errors=True)

    def out(self, name: str) -> str:
        return os.path.join(tempfile.mkdtemp(dir=self.dir), name)

    def test_pdf_to_word_rebuilds_paragraphs_styles_and_tables(self) -> None:
        dst = self.out("r.docx")
        engines.pdf_to_docx(self.pdf, dst)
        d = Document(dst)
        texts = [p.text for p in d.paragraphs if p.text.strip()]
        self.assertEqual(texts[0], "Quarterly report")
        joined = next(t for t in texts if t.startswith("The first paragraph"))
        self.assertIn("wrap across more than one line, so the lines must be joined again.", joined)  # one paragraph
        self.assertTrue(any(r.bold and "bold words" in r.text for p in d.paragraphs for r in p.runs))
        self.assertIn("Page 2 text", texts)
        self.assertEqual(len(d.tables), 1)
        self.assertEqual([c.text for c in d.tables[0].rows[1].cells], ["Apples", "12", "1,250.50"])

    def test_pdf_to_excel_types_numbers_and_never_writes_formulas(self) -> None:
        dst = self.out("r.xlsx")
        engines.pdf_to_xlsx(self.pdf, dst)
        wb = load_workbook(dst)
        self.assertEqual(wb.sheetnames, ["Page 1", "Page 2"])
        cells = {c.value: c for row in wb["Page 1"].iter_rows() for c in row if c.value is not None}
        self.assertEqual(cells[12].data_type, "n")
        self.assertIn(1250.5, cells)
        # Text that looks like a formula must stay text (no formula injection into the user's spreadsheet).
        risky = [c for c in cells.values() if isinstance(c.value, str) and c.value.startswith("=")]
        self.assertTrue(risky)
        for c in risky:
            self.assertEqual(c.data_type, "s", c.value)
        with zipfile.ZipFile(dst) as z:
            sheet = z.read("xl/worksheets/sheet1.xml").decode()
        self.assertNotIn("<f>", sheet)

    def test_pdf_to_powerpoint_one_slide_per_page_with_text_in_notes(self) -> None:
        dst = self.out("r.pptx")
        engines.pdf_to_pptx(self.pdf, dst)
        p = Presentation(dst)
        self.assertEqual(len(p.slides), 2)
        self.assertEqual(len(p.slides[0].shapes), 1)
        self.assertIn("Quarterly report", p.slides[0].notes_slide.notes_text_frame.text)
        self.assertNotIn("_x000D_", p.slides[0].notes_slide.notes_text_frame.text)

    def test_office_to_pdf_for_word_excel_and_powerpoint(self) -> None:
        wb = Workbook()
        wb.active["A1"] = "Budget"
        wb.active["B2"] = 42
        xlsx = os.path.join(tempfile.mkdtemp(dir=self.dir), "input.xlsx")
        wb.save(xlsx)
        pres = Presentation()
        s = pres.slides.add_slide(pres.slide_layouts[0])
        s.shapes.title.text = "Launch plan"
        pptx = os.path.join(tempfile.mkdtemp(dir=self.dir), "input.pptx")
        pres.save(pptx)
        docx = os.path.join(tempfile.mkdtemp(dir=self.dir), "input.docx")
        report_docx(docx)
        import pypdfium2 as pdfium

        for src, ext, needle in [(docx, "docx", "Quarterly report"), (xlsx, "xlsx", "Budget"), (pptx, "pptx", "Launch plan")]:
            dst = self.out("out.pdf")
            engines.office_to_pdf(src, dst, ext, 120)
            doc = pdfium.PdfDocument(dst)
            self.assertIn(needle, doc[0].get_textpage().get_text_range(), ext)
            doc.close()

    def test_office_to_pdf_does_not_fetch_linked_remote_content(self) -> None:
        hits: list[str] = []

        class H(http.server.BaseHTTPRequestHandler):
            def do_GET(self) -> None:
                hits.append(self.path)
                self.send_response(404)
                self.end_headers()

            def log_message(self, *_a) -> None:
                pass

        srv = http.server.HTTPServer(("127.0.0.1", 0), H)
        threading.Thread(target=srv.serve_forever, daemon=True).start()
        try:
            src = os.path.join(tempfile.mkdtemp(dir=self.dir), "input.docx")
            linked_image_docx(src, f"http://127.0.0.1:{srv.server_port}/tracker.png")
            dst = self.out("out.pdf")
            engines.office_to_pdf(src, dst, "docx", 120)
            self.assertTrue(os.path.getsize(dst) > 0)
        finally:
            srv.shutdown()
            srv.server_close()
        self.assertEqual(hits, [], "LibreOffice fetched remote content linked from the document")

    def test_refuses_password_protected_pdfs(self) -> None:
        for fn in (engines.pdf_to_docx, engines.pdf_to_xlsx, engines.pdf_to_pptx):
            with self.assertRaisesRegex(UserError, "password-protected"):
                fn(self.encrypted, self.out("x"))

    def test_refuses_damaged_files_and_too_many_pages(self) -> None:
        bad = os.path.join(self.dir, "bad.pdf")
        with open(bad, "wb") as f:
            f.write(b"%PDF-1.7\n this is not a pdf")
        with self.assertRaisesRegex(UserError, "could not be read"):
            engines.pdf_to_docx(bad, self.out("x.docx"))
        old = engines.MAX_PAGES
        engines.MAX_PAGES = 1
        try:
            with self.assertRaisesRegex(UserError, "Convert up to 1 pages"):
                engines.pdf_to_xlsx(self.pdf, self.out("x.xlsx"))
        finally:
            engines.MAX_PAGES = old
        with self.assertRaisesRegex(UserError, "cannot be converted"):
            engines.office_to_pdf(bad, self.out("x.pdf"), "exe", 10)


def linked_image_docx(path: str, url: str) -> None:
    """A Word file whose only picture is linked (not embedded) from `url`."""
    base = io.BytesIO()
    d = Document()
    d.add_paragraph("Linked picture below")
    d.save(base)
    with zipfile.ZipFile(base) as z:
        files = {n: z.read(n) for n in z.namelist()}
    rel = f'<Relationship Id="rIdX" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="{url}" TargetMode="External"/>'
    files["word/_rels/document.xml.rels"] = files["word/_rels/document.xml.rels"].decode().replace("</Relationships>", rel + "</Relationships>").encode()
    pic = (
        '<w:p><w:r><w:drawing><wp:inline xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">'
        '<wp:extent cx="914400" cy="914400"/><wp:docPr id="9" name="p"/>'
        '<a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
        '<pic:pic xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture"><pic:nvPicPr><pic:cNvPr id="9" name="p"/><pic:cNvPicPr/></pic:nvPicPr>'
        '<pic:blipFill><a:blip xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:link="rIdX"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>'
        '<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="914400" cy="914400"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr>'
        "</pic:pic></a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p>"
    )
    files["word/document.xml"] = files["word/document.xml"].decode().replace("<w:sectPr", pic + "<w:sectPr", 1).encode()
    with zipfile.ZipFile(path, "w") as z:
        for n, b in files.items():
            z.writestr(n, b)


class ServerTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        env = {"CONVERTER_TOKEN": TOKEN, "CONVERTER_MAX_BYTES": str(2 * 1024 * 1024), "CONVERTER_QUEUE_WAIT": "0"}
        old = dict(os.environ)
        os.environ.update(env)
        try:
            cls.settings = server.Settings()
        finally:
            os.environ.clear()
            os.environ.update(old)
        cls.slots = threading.BoundedSemaphore(1)
        cls.srv = server.make_server("127.0.0.1", 0, server.make_handler(cls.settings, cls.slots))
        threading.Thread(target=cls.srv.serve_forever, daemon=True).start()
        cls.dir = tempfile.mkdtemp()
        if HAS_SOFFICE:
            report_docx(os.path.join(cls.dir, "r.docx"))
            with open(soffice_pdf(os.path.join(cls.dir, "r.docx"), cls.dir), "rb") as fixture:
                cls.pdf = fixture.read()

    @classmethod
    def tearDownClass(cls) -> None:
        cls.srv.shutdown()
        cls.srv.server_close()
        shutil.rmtree(cls.dir, ignore_errors=True)

    def post(self, path: str, body: bytes, token: str | None = TOKEN, headers: dict | None = None):
        c = http.client.HTTPConnection("127.0.0.1", self.srv.server_address[1], timeout=180)
        h = {"content-type": "application/octet-stream", **(headers or {})}
        if token is not None:
            h["authorization"] = f"Bearer {token}"
        c.request("POST", path, body=body, headers=h)
        r = c.getresponse()
        data = r.read()
        c.close()
        return r.status, r.getheader("content-type"), data

    def test_settings_require_a_strong_token(self) -> None:
        old = dict(os.environ)
        try:
            os.environ["CONVERTER_TOKEN"] = "short"
            with self.assertRaises(SystemExit):
                server.Settings()
            os.environ.pop("CONVERTER_TOKEN")
            with self.assertRaises(SystemExit):
                server.Settings()
        finally:
            os.environ.clear()
            os.environ.update(old)

    def test_health_identifies_converter_and_checks_dependencies(self) -> None:
        for missing, expected in [(False, 200), (True, 503)]:
            with patch.object(server.shutil, "which", return_value=None if missing else "/usr/bin/soffice"):
                c = http.client.HTTPConnection("127.0.0.1", self.srv.server_address[1])
                c.request("GET", "/health")
                res = c.getresponse()
                self.assertEqual(res.status, expected)
                self.assertEqual(res.getheader("x-rokela-service"), "converter")
                self.assertEqual(json.loads(res.read()), {"ok": not missing, "service": "rokela-converter"})
                c.close()

    @unittest.skipUnless(HAS_SOFFICE, "LibreOffice is not installed")
    def test_all_conversions_through_the_resource_limited_worker(self) -> None:
        # Engine tests run in-process. These exercise job.py with the same user,
        # environment filtering and resource limits used by production HTTP requests.
        document = io.BytesIO()
        d = Document()
        d.add_paragraph("Worker conversion check")
        d.save(document)
        for target, ext, data in [("pdf", "docx", document.getvalue()),
                                  ("xlsx", "pdf", self.pdf), ("pptx", "pdf", self.pdf)]:
            status, ctype, body = self.post(f"/v1/convert/{target}?from={ext}", data, headers={"x-request-id": "test-worker"})
            self.assertEqual(status, 200, (target, body[:300]))
            self.assertEqual(ctype, server.OUTPUT_TYPES[target])
            self.assertTrue(body.startswith(b"%PDF-") if target == "pdf" else zipfile.is_zipfile(io.BytesIO(body)))

    def test_requires_the_shared_token(self) -> None:
        for tok in (None, "", "x" * 40, TOKEN + "x", TOKEN[:-1]):
            status, _, _ = self.post("/v1/convert/docx?from=pdf", b"%PDF-1.7\n", token=tok)
            self.assertEqual(status, 401, tok)

    def test_rejects_unknown_routes_wrong_types_and_mismatched_content(self) -> None:
        self.assertEqual(self.post("/v1/convert/exe?from=pdf", b"%PDF-1.7")[0], 404)
        self.assertEqual(self.post("/v1/other/docx?from=pdf", b"%PDF-1.7")[0], 404)
        self.assertEqual(self.post("/v1/convert/docx?from=docx", b"PK\x03\x04")[0], 400)
        self.assertEqual(self.post("/v1/convert/pdf?from=pdf", b"%PDF-1.7")[0], 400)
        self.assertEqual(self.post("/v1/convert/pdf?from=html", b"<html>")[0], 400)
        status, _, body = self.post("/v1/convert/pdf?from=docx", b"%PDF-1.7 pretending to be a docx")
        self.assertEqual((status, json.loads(body)["code"]), (400, "unsupported"))
        self.assertEqual(self.post("/v1/convert/docx?from=pdf", b"PK\x03\x04 not a pdf")[0], 400)

    def test_enforces_the_size_limit(self) -> None:
        # Refusal happens before the body is read. Sending a huge body races the
        # server closing the socket and can raise BrokenPipeError in http.client.
        status, _, body = self.post("/v1/convert/docx?from=pdf", b"%PDF-1.7\n", headers={"content-length": str(2 * 1024 * 1024 + 1)})
        self.assertEqual((status, json.loads(body)["code"]), (413, "too_large"))

    def test_says_busy_instead_of_queueing_without_limit(self) -> None:
        self.slots.acquire()
        try:
            status, _, body = self.post("/v1/convert/docx?from=pdf", b"%PDF-1.7\n")
        finally:
            self.slots.release()
        self.assertEqual((status, json.loads(body)["code"]), (429, "busy"))

    @unittest.skipUnless(HAS_SOFFICE, "LibreOffice is not installed")
    def test_converts_and_cleans_up(self) -> None:
        before = set(os.listdir(tempfile.gettempdir()))
        status, ctype, body = self.post("/v1/convert/docx?from=pdf", self.pdf)
        self.assertEqual(status, 200)
        self.assertEqual(ctype, server.OUTPUT_TYPES["docx"])
        self.assertIn("Quarterly report", "\n".join(p.text for p in Document(io.BytesIO(body)).paragraphs))
        leftovers = [n for n in set(os.listdir(tempfile.gettempdir())) - before if n.startswith("job-")]
        self.assertEqual(leftovers, [])

    def test_damaged_file_is_a_clear_error(self) -> None:
        status, _, body = self.post("/v1/convert/xlsx?from=pdf", b"%PDF-1.7\n garbage")
        self.assertEqual(status, 422)
        self.assertIn("could not be read", json.loads(body)["message"])


if __name__ == "__main__":
    unittest.main()
